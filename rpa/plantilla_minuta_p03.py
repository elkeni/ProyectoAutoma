"""Ensamblador determinista del borrador de minuta de compraventa."""

from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from typing import Literal

from docx import Document
from pydantic import BaseModel, ConfigDict, Field


class DatosMinuta(BaseModel):
    """Datos ya validados por P01/P02; el ensamblador no completa hechos."""

    model_config = ConfigDict(extra="forbid", frozen=True, str_strip_whitespace=True)

    expediente_id: str = Field(min_length=1, max_length=128)
    vendedor: str = Field(min_length=2, max_length=200)
    comprador: str = Field(min_length=2, max_length=200)
    inmueble: str = Field(min_length=3, max_length=500)
    partida_registral: str = Field(min_length=3, max_length=40)
    valor_comercial: str = Field(min_length=1, max_length=80)
    forma_pago: str = Field(min_length=3, max_length=200)
    obligaciones_adicionales: tuple[str, ...] = ()
    referencias: tuple[str, ...] = ()


class BorradorMinuta(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    expediente_id: str
    formato: Literal["markdown", "docx"]
    contenido_markdown: str
    contenido_docx: bytes | None = None
    hash_sha256: str
    clausulas: tuple[str, ...]
    nodo_bpmn_id: Literal["p3_borrador"] = "p3_borrador"


def _clausulas(datos: DatosMinuta) -> list[tuple[str, str]]:
    adicionales = (
        "\n".join(f"- {obligacion}" for obligacion in datos.obligaciones_adicionales)
        or "- No se han consignado obligaciones adicionales."
    )
    referencias = ", ".join(datos.referencias) or "Sin referencias documentales declaradas"
    return [
        (
            "PRIMERA: PARTES",
            f"Intervienen como parte vendedora {datos.vendedor} y como parte "
            f"compradora {datos.comprador}.",
        ),
        (
            "SEGUNDA: OBJETO",
            f"La parte vendedora transfiere el inmueble {datos.inmueble}, inscrito "
            f"en la partida registral {datos.partida_registral}.",
        ),
        (
            "TERCERA: PRECIO Y FORMA DE PAGO",
            f"El precio acordado es {datos.valor_comercial}. La forma de pago es "
            f"{datos.forma_pago}.",
        ),
        (
            "CUARTA: DECLARACIONES Y SANEAMIENTO",
            "Las partes declaran que los datos y recaudos serán verificados en la "
            "revisión jurídica y que rigen las obligaciones legales de saneamiento.",
        ),
        ("QUINTA: OBLIGACIONES ADICIONALES", adicionales),
        (
            "SEXTA: ELEVACIÓN A ESCRITURA PÚBLICA",
            "Las partes se obligan a suscribir los instrumentos necesarios para la "
            "formalización notarial una vez aprobada y firmada la minuta.",
        ),
        ("REFERENCIAS DE TRAZABILIDAD", referencias),
    ]


def generar_markdown(datos: DatosMinuta) -> tuple[str, tuple[str, ...]]:
    clausulas = _clausulas(datos)
    lineas = [
        "# BORRADOR DE MINUTA DE COMPRAVENTA",
        "",
        f"**Expediente:** {datos.expediente_id}",
        "",
        "> BORRADOR SUJETO A REVISIÓN JURÍDICA Y FIRMA HUMANA.",
        "",
    ]
    for titulo, cuerpo in clausulas:
        lineas.extend((f"## {titulo}", "", cuerpo, ""))
    contenido = "\n".join(lineas).rstrip() + "\n"
    return contenido, tuple(titulo for titulo, _ in clausulas)


def generar_docx(datos: DatosMinuta) -> bytes:
    documento = Document()
    documento.add_heading("BORRADOR DE MINUTA DE COMPRAVENTA", level=0)
    documento.add_paragraph(f"Expediente: {datos.expediente_id}")
    documento.add_paragraph("BORRADOR SUJETO A REVISIÓN JURÍDICA Y FIRMA HUMANA.")
    for titulo, cuerpo in _clausulas(datos):
        documento.add_heading(titulo, level=1)
        documento.add_paragraph(cuerpo)
    salida = BytesIO()
    documento.save(salida)
    return salida.getvalue()


def ensamblar_minuta(
    datos: DatosMinuta,
    *,
    formato: Literal["markdown", "docx"] = "markdown",
    ruta_salida: str | Path | None = None,
) -> BorradorMinuta:
    """Ejecuta ``p3_borrador`` sin inferencias ni cláusulas inventadas."""

    markdown, nombres = generar_markdown(datos)
    binario = generar_docx(datos) if formato == "docx" else None
    material = binario if binario is not None else markdown.encode("utf-8")
    if ruta_salida is not None:
        ruta = Path(ruta_salida)
        if formato == "docx":
            ruta.write_bytes(material)
        else:
            ruta.write_text(markdown, encoding="utf-8")
    return BorradorMinuta(
        expediente_id=datos.expediente_id,
        formato=formato,
        contenido_markdown=markdown,
        contenido_docx=binario,
        hash_sha256=hashlib.sha256(material).hexdigest(),
        clausulas=nombres,
    )


__all__ = [
    "BorradorMinuta",
    "DatosMinuta",
    "ensamblar_minuta",
    "generar_docx",
    "generar_markdown",
]
